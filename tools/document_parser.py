"""
document_parser.py — 文档解析与智能切片工具
"""
import os
import subprocess
import tempfile
from pathlib import Path

# 确保安装了依赖：pip install pdfplumber python-docx
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None

# OCR fallback（扫描图片 PDF）
try:
    from pdfminer.high_level import extract_text as _pdfminer_extract
    _HAS_PDFMINER = True
except ImportError:
    _HAS_PDFMINER = False

try:
    import pytesseract
    from pdf2image import convert_from_path
    _HAS_OCR = True
except ImportError:
    _HAS_OCR = False


def _find_soffice() -> str | None:
    """查找 LibreOffice soffice 可执行文件"""
    import shutil
    # 1. PATH 中直接可用
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    # 2. Windows 常见安装路径
    if os.name == "nt":
        for p in [r"C:\Program Files\LibreOffice\program\soffice.exe",
                   r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
            if os.path.exists(p):
                return p
    # 3. Linux 常见安装路径
    for p in ["/usr/bin/soffice", "/usr/bin/libreoffice", "/usr/lib/libreoffice/program/soffice"]:
        if os.path.exists(p):
            return p
    return None


def extract_text(file_path: str) -> str:
    """自动识别后缀，提取纯文本"""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"❌ 找不到文件：{file_path}")

    ext = path.suffix.lower()
    text = ""

    if ext == ".pdf":
        if not pdfplumber:
            raise ImportError("请先安装 pdfplumber: pip install pdfplumber")
        print(f"  [解析] 打开 PDF: {path.name} ({path.stat().st_size} bytes)", flush=True)
        with pdfplumber.open(path) as pdf:
            print(f"  [解析] PDF 共 {len(pdf.pages)} 页", flush=True)
            empty_pages = 0
            for i, page in enumerate(pdf.pages):
                extracted = page.extract_text()
                chars = len(extracted) if extracted else 0
                if chars > 0:
                    text += extracted + "\n\n"
                else:
                    empty_pages += 1
            print(f"  [解析] pdfplumber 提取: {len(text)} 字符 ({empty_pages}/{len(pdf.pages)} 页为空)", flush=True)

        # ── Fallback 1: pdfminer（不同的文本提取引擎）──
        if not text.strip() and _HAS_PDFMINER:
            print("  [解析] pdfplumber 提取为空，尝试 pdfminer...", flush=True)
            try:
                text = _pdfminer_extract(str(path))
                print(f"  [解析] pdfminer 提取: {len(text)} 字符", flush=True)
            except Exception as e:
                print(f"  [解析] pdfminer 失败: {e}", flush=True)

        # ── Fallback 2: OCR（扫描图片 PDF）──
        if not text.strip() and _HAS_OCR:
            print("  [解析] 文本提取均为空，启动 OCR 识别扫描图片...", flush=True)
            try:
                images = convert_from_path(str(path))
                ocr_parts = []
                for i, img in enumerate(images):
                    ocr_text = pytesseract.image_to_string(img, lang="eng+chi_sim")
                    if ocr_text.strip():
                        ocr_parts.append(ocr_text)
                    print(f"  [解析] OCR 第 {i+1} 页: {len(ocr_text)} 字符", flush=True)
                text = "\n\n".join(ocr_parts)
                print(f"  [解析] OCR 总计: {len(text)} 字符", flush=True)
            except Exception as e:
                print(f"  [解析] OCR 失败: {e}", flush=True)

        # ── 诊断提示 ──
        if not text.strip():
            msg_parts = ["⚠️ PDF 提取文本为空！"]
            if not _HAS_OCR:
                msg_parts.append("这是扫描图片 PDF，需要安装 OCR 支持：")
                msg_parts.append("  pip install pytesseract pdf2image")
                msg_parts.append("  并安装 Tesseract: https://tesseract-ocr.github.io/tessdoc/Installation.html")
            else:
                msg_parts.append("OCR 也未能识别，PDF 可能损坏或为纯图片格式。")
            print("\n".join(msg_parts), flush=True)

    elif ext == ".doc":
        print(f"  [解析] 检测到旧版 .doc 格式，尝试通过 LibreOffice 转换为 .docx...", flush=True)
        tmpdir = tempfile.mkdtemp(prefix="doc_conv_")
        try:
            # 查找 LibreOffice / soffice
            soffice = _find_soffice()
            if soffice is None:
                raise RuntimeError(
                    "未找到 LibreOffice/soffice。服务器请执行: apt install libreoffice-writer；"
                    "Windows 请安装 LibreOffice 并确保 soffice.exe 在 PATH 中。"
                )
            # 复制文件到临时目录（避免路径含空格/中文导致 LibreOffice 报错）
            safe_src = os.path.join(tmpdir, "input.doc")
            import shutil
            shutil.copy2(str(path), safe_src)
            # 调用 LibreOffice 无头转换
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmpdir, safe_src],
                capture_output=True, timeout=60,
            )
            converted = os.path.join(tmpdir, "input.docx")
            if not os.path.exists(converted):
                raise RuntimeError("LibreOffice 转换 .doc → .docx 失败，未生成输出文件")
            print(f"  [解析] 转换成功，开始解析 docx...", flush=True)
            if not docx:
                raise ImportError("请先安装 python-docx: pip install python-docx")
            doc_obj = docx.Document(converted)
            text = "\n".join([para.text for para in doc_obj.paragraphs if para.text.strip()])
            print(f"  [解析] .doc 提取: {len(text)} 字符", flush=True)
        finally:
            import shutil as _sh
            _sh.rmtree(tmpdir, ignore_errors=True)

    elif ext == ".docx":
        if not docx:
            raise ImportError("请先安装 python-docx: pip install python-docx")
        doc = docx.Document(path)
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    elif ext in [".md", ".txt"]:
        text = path.read_text(encoding="utf-8")

    else:
        raise ValueError(f"⚠️ 暂不支持该文件格式解析：{ext}")

    return text


def chunk_text(text: str, chunk_size: int = 4000, overlap: int = 300) -> list[str]:
    """
    智能切片：带重叠的滑动窗口算法
    chunk_size: 每个切片的最大字符数
    overlap: 切片之间的重叠字符数，防止概念被生硬切断
    """
    if not text:
        return []

    chunks = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = start + chunk_size

        # 如果不是最后一块，尽量在换行符处切断，保持段落完整性
        if end < text_length:
            # 往回找最近的换行符
            last_newline = text.rfind('\n', start, end)
            if last_newline != -1 and last_newline > start + (chunk_size // 2):
                end = last_newline + 1  # 切在换行符之后

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        # 下一个 chunk 的起点回退 overlap 的长度
        start = end - overlap

    return chunks